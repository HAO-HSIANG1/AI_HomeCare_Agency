"""Pure operational schedule/route views. Rendering and export never call map APIs."""
CAREFLOW_BUILD = '2026.10.18-v19'
from io import BytesIO
import hashlib
import json
import math

import pandas as pd


def text(value):
    if value is None or pd.isna(value):
        return ''
    return str(value).strip()


def number(value):
    try:
        value = float(value)
        return value if math.isfinite(value) else None
    except (ValueError, TypeError):
        return None


def day_string(value):
    value = text(value)
    if not value or value == '單日':
        return value
    if len(value) >= 10 and value[4:5] == '-' and value[7:8] == '-':
        from datetime import datetime
        try:
            return datetime.fromisoformat(value).date().isoformat()
        except ValueError:
            pass
    parsed = pd.to_datetime(value, errors='coerce')
    return parsed.strftime('%Y-%m-%d') if pd.notna(parsed) else value


def clock(value):
    value = text(value)
    if not value:
        return ''
    if len(value) == 5 and value[2] == ':' and value[:2].isdigit() and value[3:].isdigit():
        if int(value[:2]) < 24 and int(value[3:]) < 60:
            return value
    parsed = pd.to_datetime(value, errors='coerce')
    return parsed.strftime('%H:%M') if pd.notna(parsed) else value


def _taipei_stamp(day, hhmm, plus_minutes=0.0):
    """'YYYY-MM-DD' + 'HH:MM' (+分鐘) → 台北時間 ISO 字串（台灣沒有日光節約，固定 +08:00）。
    與 pd.Timestamp(...).tz_localize('Asia/Taipei').isoformat() 結果相同，但不必每筆交給 pandas 解析。"""
    from datetime import datetime, timedelta
    try:
        ts = datetime.strptime(f'{day[:10]} {hhmm}', '%Y-%m-%d %H:%M')
    except ValueError:
        ts = pd.to_datetime(day + ' ' + hhmm, errors='coerce')
        if pd.isna(ts):
            return ''
        ts = ts.to_pydatetime()
    ts += timedelta(minutes=float(plus_minutes or 0.0))
    return ts.isoformat() + '+08:00'


def describe_routes(assignments, tasks, caregivers, buffer_mins=15):
    """Describe home→first visit→subsequent visits, resetting at each worker/day."""
    columns = ['任務ID', '派單居服員', '交通日期', '交通方式', '出發地類型', '起點案號',
               '前一任務ID', '交通起點', '交通終點', '交通路段', '出發地緯度', '出發地經度',
               '終點緯度', '終點經度', '前案結束時間', '交通查詢方式', '交通查詢時間',
               '路段緩衝(分)', '交通資料簽章']
    if assignments.empty:
        return pd.DataFrame(columns=columns)
    task_map = tasks.drop_duplicates('任務ID', keep='last').set_index('任務ID').to_dict('index')
    cg_map = caregivers.drop_duplicates('居服員ID', keep='last').set_index('居服員ID').to_dict('index')
    groups = {}
    for row in assignments.drop_duplicates('任務ID', keep='last').to_dict('records'):
        task_id, cg_id = text(row.get('任務ID')), text(row.get('派單居服員'))
        if cg_id in ('', '未指派') or task_id not in task_map:
            continue
        task = task_map[task_id]
        day = day_string(task.get('日期')) or day_string(row.get('排班日期')) or '單日'
        groups.setdefault((cg_id, day), []).append(task_id)
    rows = []
    stamp_cache = {}
    for (cg_id, day), task_ids in sorted(groups.items()):
        cg = cg_map.get(cg_id, {})
        mode = text(cg.get('常用交通工具')) or text(cg.get('交通工具')) or '未提供'
        if any(word in mode for word in ('公車', '捷運', '大眾')):
            mode = '大眾運輸'
        prev_id = ''
        for task_id in sorted(task_ids, key=lambda key: (clock(task_map[key].get('時間窗_開始')), key)):
            task = task_map[task_id]
            prev = task_map.get(prev_id, {})
            first = not prev_id
            origin_id = text(prev.get('案家ID'))
            origin = '居服員住家／服務起點' if first else f'前案｜{origin_id}（{prev_id}）'
            destination = f"{text(task.get('案家ID'))}（{task_id}）"
            lat1 = number(cg.get('服務起點_緯度(家)')) if first else number(prev.get('服務地點_緯度'))
            lon1 = number(cg.get('服務起點_經度(家)')) if first else number(prev.get('服務地點_經度'))
            lat2, lon2 = number(task.get('服務地點_緯度')), number(task.get('服務地點_經度'))
            buffer = 0.0 if first or origin_id == text(task.get('案家ID')) else float(buffer_mins)
            query_clock = clock(task.get('時間窗_開始')) if first else clock(prev.get('時間窗_結束'))
            query_time = ''
            if day != '單日' and query_clock:
                stamp = (day, query_clock, 0.0 if first else buffer)
                if stamp not in stamp_cache:
                    stamp_cache[stamp] = _taipei_stamp(day, query_clock, 0.0 if first else buffer)
                query_time = stamp_cache[stamp]
            fingerprint = [task_id, cg_id, day, clock(task.get('時間窗_開始')),
                           clock(task.get('時間窗_結束')), mode, prev_id, query_time,
                           lat1, lon1, lat2, lon2, buffer]
            rows.append(dict(zip(columns, [task_id, cg_id, day, mode,
                '住家／服務起點' if first else '前一案家', origin_id, prev_id,
                origin, destination, origin + ' → ' + destination, lat1, lon1, lat2, lon2,
                clock(prev.get('時間窗_結束')), '到達' if first else '出發', query_time, buffer,
                hashlib.sha256(json.dumps(fingerprint, ensure_ascii=False).encode()).hexdigest()])))
            prev_id = task_id
    return pd.DataFrame(rows, columns=columns)


MAIN_COLUMNS = ['日期', '星期', '開始時間', '結束時間', '案號', '派單居服員', '服務項目',
                '服務分鐘', '出發地', '交通方式', '交通時間', '路段緩衝(分)',
                '含緩衝交通時間(分)', '交通計算狀態', '任務ID']


def build_schedule_output(tasks, assignments, caregivers, routes, metadata=None, buffer_mins=15):
    """Use only travel calculated for the current worker, predecessor, and time."""
    sequence = describe_routes(assignments, tasks, caregivers, buffer_mins)
    sequence_map = sequence.set_index('任務ID').to_dict('index') if not sequence.empty else {}
    route_map = routes.drop_duplicates('任務ID', keep='last').set_index('任務ID').to_dict('index') if not routes.empty else {}
    meta_map = metadata.drop_duplicates('任務ID', keep='last').set_index('任務ID').to_dict('index') if metadata is not None and not metadata.empty else {}
    rows = []
    for _, task in tasks.drop_duplicates('任務ID', keep='last').iterrows():
        task_id = text(task['任務ID'])
        row = dict(task)
        row.update(meta_map.get(task_id, {}))
        seq = sequence_map.get(task_id)
        route = route_map.get(task_id, {})
        row.update({'任務ID': task_id, '案號': task.get('案家ID', task.get('案號')),
                    '日期': day_string(task.get('日期', '')), '星期': task.get('星期', ''),
                    '開始時間': clock(task.get('時間窗_開始', task.get('開始時間'))),
                    '結束時間': clock(task.get('時間窗_結束', task.get('結束時間')))})
        minutes = number(task.get('服務歷時(分鐘)', task.get('服務分鐘')))
        row['服務分鐘'] = minutes
        row['服務時數'] = round(minutes / 60, 4) if minutes is not None else None
        service_parts = []
        for col in sorted(task.index):
            if col.startswith('Service_Code_') and text(task[col]):
                suffix = col.rsplit('_', 1)[1]
                units = text(task.get('Units_' + suffix))
                service_parts.append(text(task[col]) + (f'({units})' if units else ''))
        row['服務項目'] = '、'.join(service_parts) or text(task.get('服務項目'))
        row['派單居服員'] = seq['派單居服員'] if seq else '未指派'
        row['最終派單居服員'] = row['派單居服員']
        row['交通方式'] = seq['交通方式'] if seq else ''
        row['出發地'] = seq['交通起點'] if seq else ''
        row['交通時間'] = None  # The original weekly-plan zero is a placeholder.
        row['預估車程(分)'] = None
        row['含緩衝交通時間(分)'] = None
        row['路段緩衝(分)'] = seq['路段緩衝(分)'] if seq else None
        row['交通計算狀態'] = '已派單，交通待計算' if seq else '未派單，交通尚未計算'
        row['交通估算來源'] = ''
        row['轉場時段檢查'] = ''
        if seq:
            row.update(seq)
            valid = (text(route.get('派單居服員')) == seq['派單居服員'] and
                     text(route.get('交通資料簽章')) == seq['交通資料簽章'])
            if valid:
                travel = number(route.get('預估車程(分)'))
                row['交通時間'] = travel
                row['預估車程(分)'] = travel
                row['含緩衝交通時間(分)'] = number(route.get('含緩衝交通時間(分)'))
                for col in ('交通計算狀態', '交通估算來源', '交通查詢時間', '轉場時段檢查'):
                    row[col] = route.get(col, '')
            elif route:
                row['交通計算狀態'] = '派單或時段已變更，交通待重算'
        rows.append(row)
    out = pd.DataFrame(rows)
    weekly = ['狀態', '費用類型', '頻率', '生效日', '到期日', '服務時數']
    for col in MAIN_COLUMNS + weekly:
        if col not in out:
            out[col] = None
    order = MAIN_COLUMNS + weekly + [col for col in out if col not in MAIN_COLUMNS + weekly]
    return out[order].sort_values(['日期', '開始時間', '派單居服員', '任務ID'], na_position='last').reset_index(drop=True)


def schedule_excel_bytes(schedule, daily_payroll=None, monthly_payroll=None, primary_map=None, diagnostics=None):
    """Export familiar weekly-plan columns plus an auditable per-leg worksheet."""
    route_columns = ['日期', '任務ID', '案號', '派單居服員', '開始時間', '結束時間', '交通方式',
                     '出發地類型', '出發地', '起點案號', '前一任務ID', '前案結束時間',
                     '出發地緯度', '出發地經度', '終點緯度', '終點經度', '交通查詢方式',
                     '交通查詢時間', '交通時間', '路段緩衝(分)', '含緩衝交通時間(分)',
                     '交通估算來源', '交通計算狀態', '轉場時段檢查', '交通資料簽章']
    notes = pd.DataFrame([
        ('交通時間', '系統估計分鐘，不含轉場緩衝；原週服務計畫的0是待計算占位值。空白表示未知，不能當作0分鐘。'),
        ('服務時數', '服務分鐘÷60的小數小時；例如75分鐘=1.25小時，不使用原檔HH.MM式數字。'),
        ('出發地', '每人每天第一案由員工資料的住家／服務起點出發；其他由當日有效派單順序的前一案家出發。非實際GPS軌跡。'),
        ('交通方式', '依員工資料選機車或大眾運輸。大眾運輸可能含步行；此欄不能證明實際搭乘公車／捷運。'),
        ('交通查詢方式', '首案以服務開始作為到達條件；後續案以前案結束+轉場緩衝作為出發條件。OSRM不提供即時交通預測。'),
        ('路段緩衝(分)', '獨立預留停車、步行、門禁等時間；首案顯示0，不代表出門與門禁不用時間。'),
        ('交通資料簽章', '用於防止改派／改期／插單後沿用舊路段交通。交通估算仍須人工核對或實測。'),
        ('班表範圍', '含未指派與臨時插單。逐段交通不含最後一案返家。指派結果為排班建議。'),
        ('固定居服員ID', '每案家跨所有星期與時段固定同一人；已登錄且與服務重疊的請假才開放代班，代班不覆寫固定人選。'),
        ('SA10計費', '全部為自費，一單位一小時，採設定之自費時薪估價；不計入公費給付。若來源單位與時段不符，保留來源時段並列提醒。'),
        ('服務基本工資', '實際排定服務分鐘÷60×設定時薪（預設210）。與服務收費無關。'),
        ('轉場工資', '同人同日依有效服務順序，每次不同案家間轉場46元；不另加210×轉場小時。首案通勤、末案返家及同案接續服務不計轉場工資。'),
        ('加班加給', '一般工作日服務加轉場與工作緩衝超過8小時，前2小時加給210×0.34，後2小時加給210×0.67；基本工資已另列，避免再付完整倍率。非休息日／例假日／國定假日完整薪資結算。'),
        ('月工時檢查', '本次有效班表逐日超過8小時的時數加總，預設上限46小時；可加員工欄位「排班期外本月加班時數」。未提供之期外出勤需另核對。'),
    ], columns=['欄位', '說明'])
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine='openpyxl') as writer:
        schedule.to_excel(writer, sheet_name='班表', index=False)
        schedule.reindex(columns=route_columns).to_excel(writer, sheet_name='逐段交通', index=False)
        if '日期' in schedule and ('案家ID' in schedule or '案號' in schedule):
            from pending_clients import pending_client_summary
            view = schedule.copy()
            for target,source in [('案家ID','案號'),('時間窗_開始','開始時間'),('時間窗_結束','結束時間')]:
                if target not in view and source in view:
                    view[target]=view[source]
            pending_client_summary(view,view,diagnostics,primary_map).to_excel(writer,sheet_name='未派案家',index=False)
        if diagnostics is not None:
            pending = diagnostics.copy()
            if '逐人排班原因' in pending:
                import json
                pending['逐人排班原因']=pending['逐人排班原因'].map(lambda v:json.dumps(v,ensure_ascii=False) if isinstance(v,dict) else v)
            pending.to_excel(writer,sheet_name='未派原因',index=False)
        notes.to_excel(writer, sheet_name='欄位說明', index=False)
        if daily_payroll is not None:
            daily_payroll.to_excel(writer, sheet_name='每日工資估算', index=False)
        if monthly_payroll is not None:
            monthly_payroll.to_excel(writer, sheet_name='每月工資估算', index=False)
        if primary_map is not None:
            pd.DataFrame([{'案家ID':client,'固定居服員ID':worker} for client,worker in sorted(primary_map.items())], columns=['案家ID','固定居服員ID']).to_excel(writer, sheet_name='固定居服員', index=False)
        from openpyxl.styles import Alignment, Font, PatternFill
        for ws in writer.book:
            ws.freeze_panes = 'E2' if ws.title == '班表' else 'A2'
            ws.auto_filter.ref = ws.dimensions
            ws.sheet_view.showGridLines = False
            for cell in ws[1]:
                cell.fill = PatternFill('solid', fgColor='176B65')
                cell.font = Font(color='FFFFFF', bold=True)
                cell.alignment = Alignment(wrap_text=True)
            ws.row_dimensions[1].height = 30
            for col in ws.columns:
                heading = str(col[0].value)
                ws.column_dimensions[col[0].column_letter].width = 48 if heading in ('出發地', '交通路段', '說明') else (27 if '時間' in heading or heading == '任務ID' else 20)
                for cell in col[1:]:
                    cell.alignment = Alignment(vertical='top', wrap_text=True)
                    if isinstance(cell.value, (int, float)):
                        cell.number_format = '0.00' if any(word in heading for word in ('交通', '時數', '緩衝')) else '0.######'
                    if isinstance(cell.value, str) and cell.value.startswith(('=', '+', '-', '@')):
                        cell.data_type = 's'
            if ws.title == '欄位說明':
                ws.column_dimensions['B'].width = 110
    return buf.getvalue()
