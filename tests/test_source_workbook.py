from io import BytesIO
from zipfile import ZipFile
from xml.dom import minidom
import pandas as pd
from openpyxl import load_workbook
from source_workbook import update_source_workbook


def test_source_export_preserves_other_sheets_and_formula_nodes():
    original=BytesIO()
    with pd.ExcelWriter(original,engine='openpyxl') as writer:
        pd.DataFrame([{'內部案號':'C','其他欄位':'保留'}]).to_excel(writer,sheet_name='個案資料',index=False)
        pd.DataFrame([{'案號':'C','開始時間':'09:00','結束時間':'10:00','服務分鐘':'=(C2-B2)*1440'}]).to_excel(writer,sheet_name='週服務計畫',index=False)
        pd.DataFrame([{'編號':'03816','取得資格':'失智症照顧服務20小時'}]).to_excel(writer,sheet_name='員工資料',index=False)
    before=original.getvalue()
    after=update_source_workbook(before,{'C':'03816'},{2:{'開始時間':'09:30','結束時間':'10:30'}})
    result=load_workbook(BytesIO(after),data_only=False)
    assert result['個案資料']['C1'].value=='歷史首選'
    assert result['個案資料']['C2'].value=='03816'
    assert result['週服務計畫']['B2'].value=='09:30'
    assert result['週服務計畫']['C2'].value=='10:30'
    assert result['週服務計畫']['D2'].value=='=(C2-B2)*1440'
    with ZipFile(BytesIO(before)) as old,ZipFile(BytesIO(after)) as new:
        assert old.namelist()==new.namelist()
        for name in old.namelist():
            if name not in ['xl/worksheets/sheet1.xml','xl/worksheets/sheet2.xml']:
                assert old.read(name)==new.read(name)


def test_existing_history_column_keeps_text_id_and_styles():
    original=BytesIO()
    with pd.ExcelWriter(original,engine='openpyxl') as writer:
        pd.DataFrame([{'內部案號':'C','歷史首選':'A'},{'內部案號':'D','歷史首選':'B'}]).to_excel(writer,sheet_name='個案資料',index=False)
    result=load_workbook(BytesIO(update_source_workbook(original.getvalue(),{'C':'03816'})))
    assert result['個案資料'].max_column==2
    assert result['個案資料']['B2'].value=='03816'
    assert result['個案資料']['B3'].value=='B'


def test_weekly_team_column_preserves_formula_cache_and_table(tmp_path):
    from openpyxl.worksheet.table import Table
    original=BytesIO()
    with pd.ExcelWriter(original,engine='openpyxl') as writer:
        pd.DataFrame([{'內部案號':'C'}]).to_excel(writer,sheet_name='個案資料',index=False)
        pd.DataFrame([{'案號':'C','開始時間':'09:00','結束時間':'10:00','服務分鐘':'=(C2-B2)*1440'}]).to_excel(writer,sheet_name='週服務計畫',index=False)
        writer.book['週服務計畫'].add_table(Table(displayName='Weekly',ref='A1:D2'))
    before=original.getvalue()
    after=update_source_workbook(before,{'C':'A'},weekly_workers={2:'03816'})
    result=load_workbook(BytesIO(after))
    sheet=result['週服務計畫']
    assert sheet['E1'].value=='週序列居服員' and sheet['E2'].value=='03816'
    assert sheet.tables['Weekly'].ref=='A1:E2'
    assert sheet['D2'].value=='=(C2-B2)*1440'
    with ZipFile(BytesIO(before)) as old,ZipFile(BytesIO(after)) as new:
        old_doc=minidom.parseString(old.read('xl/worksheets/sheet2.xml'))
        new_doc=minidom.parseString(new.read('xl/worksheets/sheet2.xml'))
        formula=lambda doc:next(c.toxml() for c in doc.getElementsByTagName('c') if c.getAttribute('r')=='D2')
        assert formula(old_doc)==formula(new_doc)
    again=load_workbook(BytesIO(update_source_workbook(after,{},weekly_workers={2:'A'})))
    assert again['週服務計畫'].max_column==5
    assert again['週服務計畫']['E2'].value=='A'
